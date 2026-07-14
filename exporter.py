"""
Funções de exportação dos dados extraídos para Excel (.xlsx) e CSV.

Tudo é gerado em memória (`io.BytesIO`), sem salvar arquivos temporários no
disco do servidor — os bytes resultantes são entregues direto para os
botões de download do Streamlit.
"""

import io

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

LARGURA_MINIMA_COLUNA = 15
LARGURA_MAXIMA_COLUNA = 50
LARGURA_COLUNA_CONFIG = 120


def results_to_dataframe(
    results: list[dict],
    abstracts: list[str],
    colunas_chave: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Junta os resultados extraídos (um dicionário por abstract) com o
    texto original em um único DataFrame. A ordem das colunas dos campos
    segue a ordem de inserção dos dicionários (mesma ordem em que o aluno
    definiu os campos), com 'abstract_original' logo após as colunas-chave
    (se houver). Se `colunas_chave` for informado, suas colunas (nomes e
    valores originais da planilha, sem alteração) são inseridas no início
    da tabela — útil para permitir que o resultado seja posteriormente
    mesclado de volta com a planilha original usando essas colunas como
    chave."""
    df = pd.DataFrame(results)
    df.insert(0, "abstract_original", abstracts)
    if colunas_chave is not None and not colunas_chave.empty:
        colunas_chave = colunas_chave.reset_index(drop=True)
        df = pd.concat([colunas_chave, df], axis=1)
    return df


def export_to_excel(df: pd.DataFrame, modelo_fonte: str, system_prompt: str) -> bytes:
    """Gera um arquivo Excel em memória com duas abas: 'Resultados' (cabeçalho
    em negrito, largura de coluna ajustada ao conteúdo e quebra de linha
    automática em células com vírgula — caso típico dos campos tipo 'lista')
    e 'Configuração' (registro do modelo Pydantic e do system prompt usados
    nessa extração, para fins de documentação/auditoria)."""
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Resultados")
        planilha = writer.sheets["Resultados"]

        for celula in planilha[1]:
            celula.font = Font(bold=True)

        for indice_coluna, nome_coluna in enumerate(df.columns, start=1):
            letra_coluna = get_column_letter(indice_coluna)
            valores_da_coluna = [str(nome_coluna)] + [str(v) for v in df[nome_coluna]]
            maior_largura = max(len(v) for v in valores_da_coluna)
            largura_final = min(max(maior_largura + 2, LARGURA_MINIMA_COLUNA), LARGURA_MAXIMA_COLUNA)
            planilha.column_dimensions[letra_coluna].width = largura_final

            for indice_linha in range(2, len(df) + 2):
                celula = planilha.cell(row=indice_linha, column=indice_coluna)
                if "," in str(celula.value):
                    celula.alignment = Alignment(wrap_text=True, vertical="top")

        aba_config = writer.book.create_sheet("Configuração")
        aba_config.column_dimensions["A"].width = LARGURA_COLUNA_CONFIG
        linha_atual = 1

        for titulo, conteudo in (
            ("Modelo Pydantic gerado", modelo_fonte),
            ("System prompt enviado à DeepSeek", system_prompt),
        ):
            celula_titulo = aba_config.cell(row=linha_atual, column=1, value=titulo)
            celula_titulo.font = Font(bold=True)
            linha_atual += 1

            for trecho in conteudo.splitlines():
                aba_config.cell(row=linha_atual, column=1, value=trecho)
                linha_atual += 1

            linha_atual += 1  # linha em branco entre as seções

    buffer.seek(0)
    return buffer.getvalue()


def export_to_csv(df: pd.DataFrame) -> bytes:
    """Gera os bytes do CSV em UTF-8 com BOM, garantindo que acentos e
    caracteres especiais apareçam corretamente ao abrir no Excel brasileiro."""
    return df.to_csv(index=False).encode("utf-8-sig")


def export_config_to_txt(modelo_fonte: str, system_prompt: str) -> bytes:
    """Gera os bytes de um .txt com o modelo Pydantic e o system prompt
    usados na extração — serve como registro de configuração para quem
    exportar em CSV, formato que não suporta múltiplas abas como o Excel."""
    secoes = (
        ("MODELO PYDANTIC GERADO", modelo_fonte),
        ("SYSTEM PROMPT ENVIADO À DEEPSEEK", system_prompt),
    )
    separador = "=" * 60
    blocos = [f"{separador}\n{titulo}\n{separador}\n{conteudo}" for titulo, conteudo in secoes]
    return "\n\n".join(blocos).encode("utf-8")
